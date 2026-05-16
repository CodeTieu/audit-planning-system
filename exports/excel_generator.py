"""
Excel workpaper generator — produces .xlsx files from workpaper form_data.
"""

import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class WorkpaperExcelGenerator:
    """Generates Excel files matching the original workpaper templates."""

    NAVY         = '1E3A5F'
    GOLD         = 'C9952A'
    YELLOW_INPUT = 'FFFFCC'   # yellow cells = auditor input

    def generate(self, workpaper) -> bytes:
        """Generate Excel file for a workpaper. Returns file bytes."""
        doc_type = workpaper.document_type.code
        method = getattr(self, f'_generate_{doc_type.lower()}', self._generate_generic)
        return method(workpaper)

    # ──────────────────────────────────────────────────────────────
    # Generic template
    # ──────────────────────────────────────────────────────────────

    def _generate_generic(self, workpaper) -> bytes:
        """Generic template for any workpaper."""
        wb = openpyxl.Workbook()

        cover = wb.active
        cover.title = 'Cover'
        self._add_cover_sheet(cover, workpaper)

        data_ws = wb.create_sheet('Form Data')
        self._add_form_data_sheet(data_ws, workpaper)

        risks_ws = wb.create_sheet('Risks Identified')
        self._add_risks_sheet(risks_ws, workpaper)

        return self._finalize(wb)

    # ──────────────────────────────────────────────────────────────
    # UE1 — Understanding the Entity: General Information
    # ──────────────────────────────────────────────────────────────

    def _generate_ue1(self, workpaper) -> bytes:
        """UE1 General Information — matches original template structure."""
        wb = openpyxl.Workbook()
        cover = wb.active
        cover.title = 'Cover'
        self._add_cover_sheet(cover, workpaper)

        ws = wb.create_sheet('Step 1 - General Information')
        fd = workpaper.form_data

        # Header
        self._header_row(ws, 1, 'UE 1 — UNDERSTANDING THE ENTITY: GENERAL INFORMATION')
        self._header_row(
            ws, 2,
            f'Engagement: {workpaper.engagement.engagement_code} | '
            f'Entity: {workpaper.engagement.entity.name} | '
            f'Period: {workpaper.engagement.audit_year}'
        )

        # Section 1
        self._section_header(ws, 4, 'SECTION 1: TYPE OF ENTITY')
        self._data_row(ws, 5, 'Entity Type',         fd.get('S1_Q1', ''))
        self._data_row(ws, 6, 'Reporting Currency',  fd.get('S1_Q2', ''))
        self._data_row(ws, 7, 'Reporting Framework', fd.get('S1_Q3', ''))

        # Section 2
        self._section_header(ws, 9, 'SECTION 2: ADDRESS & CONTACT DETAILS')
        self._data_row(ws, 10, 'Physical Address', fd.get('S2_Q1', ''))
        self._data_row(ws, 11, 'Postal Address',   fd.get('S2_Q2', ''))
        self._data_row(ws, 12, 'Telephone',        fd.get('S2_Q3', ''))
        self._data_row(ws, 13, 'Email',            fd.get('S2_Q4', ''))

        # Section 3 — Key Personnel table
        self._section_header(ws, 15, 'SECTION 3: KEY PERSONNEL')
        personnel_headers = ['Position', 'Name', 'Tenure From', 'Tenure To', 'Status', 'Notes']
        self._table_headers(ws, 16, personnel_headers)
        personnel = fd.get('key_personnel', [])
        for i, p in enumerate(personnel):
            self._table_row(ws, 17 + i, [
                p.get('position', ''), p.get('name', ''),
                p.get('tenure_from', ''), p.get('tenure_to', ''),
                p.get('status', ''),    p.get('notes', ''),
            ])

        # Section 4
        row = 18 + len(personnel)
        self._section_header(ws, row, 'SECTION 4: AUDIT ENGAGEMENT DETAILS')
        self._data_row(ws, row + 1, 'Engagement Type',  fd.get('S4_Q1', ''))
        self._data_row(ws, row + 2, 'Audit Status',     fd.get('S4_Q2', ''))
        self._data_row(ws, row + 3, 'Prior Year Opinion', fd.get('S4_Q5', ''))

        risks_ws = wb.create_sheet('Risks Identified')
        self._add_risks_sheet(risks_ws, workpaper)

        self._set_column_widths(ws, [30, 40, 15, 15, 20, 20])
        return self._finalize(wb)

    # ──────────────────────────────────────────────────────────────
    # Shared sheet builders
    # ──────────────────────────────────────────────────────────────

    def _add_cover_sheet(self, ws, workpaper):
        ws['A1'] = workpaper.document_type.name.upper()
        ws['A1'].font = Font(bold=True, size=14, color=self.NAVY)
        ws['A3'] = f'Entity: {workpaper.engagement.entity.name}'
        ws['A4'] = f'Engagement Code: {workpaper.engagement.engagement_code}'
        ws['A5'] = f'Audit Year: {workpaper.engagement.audit_year}'
        ws['A6'] = f'Status: {workpaper.status.upper()}'
        ws['A7'] = (
            f'Prepared by: {workpaper.created_by.full_name}'
            if workpaper.created_by else 'Prepared by: N/A'
        )
        ws['A8'] = (
            f'Last Updated: {workpaper.last_updated_at.strftime("%d/%m/%Y %H:%M")}'
            if workpaper.last_updated_at else 'Last Updated: N/A'
        )
        ws.column_dimensions['A'].width = 50

    def _add_form_data_sheet(self, ws, workpaper):
        ws['A1'] = 'Field'
        ws['B1'] = 'Value'
        ws['A1'].font = Font(bold=True)
        ws['B1'].font = Font(bold=True)
        for i, (key, val) in enumerate(workpaper.form_data.items(), start=2):
            ws.cell(row=i, column=1, value=str(key))
            ws.cell(
                row=i, column=2,
                value=str(val) if not isinstance(val, (list, dict)) else str(val)[:500]
            )
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 60

    def _add_risks_sheet(self, ws, workpaper):
        headers = ['Risk No.', 'Source', 'Risk Description',
                   'Severity', 'Pervasive?', 'COTABD', 'Assertions']
        self._table_headers(ws, 1, headers)
        risks = workpaper.triggered_risks.all()
        for i, risk in enumerate(risks, start=2):
            self._table_row(ws, i, [
                risk.risk_no,
                risk.source_document,
                risk.risk_description,
                risk.severity,
                'Yes' if risk.is_pervasive else 'No',
                risk.cotabd,
                ', '.join(risk.assertions or []),
            ])
        self._set_column_widths(ws, [12, 10, 60, 15, 12, 20, 30])

    # ──────────────────────────────────────────────────────────────
    # Styling helpers
    # ──────────────────────────────────────────────────────────────

    def _header_row(self, ws, row, text):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill(fill_type='solid', fgColor=self.NAVY)
        ws.merge_cells(f'A{row}:F{row}')

    def _section_header(self, ws, row, text):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, color=self.NAVY)
        cell.fill = PatternFill(fill_type='solid', fgColor='E8EDF4')
        ws.merge_cells(f'A{row}:F{row}')

    def _data_row(self, ws, row, label, value):
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = Font(bold=True)
        value_cell = ws.cell(row=row, column=2, value=str(value))
        value_cell.fill = PatternFill(fill_type='solid', fgColor=self.YELLOW_INPUT)
        ws.merge_cells(f'B{row}:F{row}')

    def _table_headers(self, ws, row, headers):
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(fill_type='solid', fgColor=self.NAVY)

    def _table_row(self, ws, row, values):
        bg = 'F8FAFC' if row % 2 == 0 else 'FFFFFF'
        fill = PatternFill(fill_type='solid', fgColor=bg)
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=str(v) if v else '')
            cell.fill = fill

    def _set_column_widths(self, ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ──────────────────────────────────────────────────────────────
    # Finalisation
    # ──────────────────────────────────────────────────────────────

    def _finalize(self, wb) -> bytes:
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()

    def protect_workbook(self, file_bytes: bytes, password: str) -> bytes:
        """Apply sheet protection to all sheets."""
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        for ws in wb.worksheets:
            ws.protection.sheet = True
            ws.protection.password = password
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()
