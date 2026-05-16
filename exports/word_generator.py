"""
Word document generator — produces the Findings Collection .docx.
"""

import io

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


class FindingsWordGenerator:
    """Generates the Findings Collection Word document."""

    def generate(self, engagement, findings) -> bytes:
        doc = Document()

        # Page setup
        section = doc.sections[0]
        section.page_width  = Inches(8.5)
        section.page_height = Inches(11)

        # Cover page
        title = doc.add_heading('AUDIT PLANNING FINDINGS', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f'Entity: {engagement.entity.name}')
        doc.add_paragraph(f'Engagement Code: {engagement.engagement_code}')
        doc.add_paragraph(f'Audit Year: {engagement.audit_year}')
        doc.add_paragraph(f'Total Findings: {len(findings)}')
        doc.add_page_break()

        # Findings summary table
        doc.add_heading('FINDINGS SUMMARY', level=1)
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = 'Table Grid'
        hdr = tbl.rows[0].cells
        hdr[0].text = 'Ref'
        hdr[1].text = 'Title'
        hdr[2].text = 'Risk Level'
        hdr[3].text = 'Status'

        for f in findings:
            row = tbl.add_row().cells
            row[0].text = f.finding_ref
            row[1].text = f.title[:60]

            linked_risks = f.risk_links.all()
            max_sev = 'Low'
            if linked_risks:
                sevs = [r.risk.severity for r in linked_risks]
                if 'High-Pervasive' in sevs:
                    max_sev = 'High-Pervasive'
                elif 'High' in sevs:
                    max_sev = 'High'
                elif 'Medium' in sevs:
                    max_sev = 'Medium'
            row[2].text = max_sev
            row[3].text = f.status.title()

        doc.add_page_break()

        # Detailed findings
        doc.add_heading('DETAILED FINDINGS', level=1)

        for i, finding in enumerate(findings, 1):
            doc.add_heading(
                f'Finding {finding.finding_ref}: {finding.title}', level=2
            )

            fields = [
                ('Criteria',                finding.criteria),
                ('Finding Body / Condition', finding.finding_body),
                ('Cause',                   finding.cause),
                ('Implication',             finding.implication),
                ('Recommendation',          finding.recommendation),
            ]

            for label, value in fields:
                if value:
                    p    = doc.add_paragraph()
                    run  = p.add_run(f'{label}: ')
                    run.bold = True
                    p.add_run(value)

            # Linked risks
            linked = finding.risk_links.select_related('risk').all()
            if linked:
                p    = doc.add_paragraph()
                run  = p.add_run('Linked Risks: ')
                run.bold = True
                p.add_run(
                    ', '.join(
                        f'{lr.risk.risk_no} ({lr.risk.severity})' for lr in linked
                    )
                )

            if i < len(findings):
                doc.add_paragraph('─' * 80)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()
